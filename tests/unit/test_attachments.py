"""只用临时文件的附件解析与范围读取测试。"""

import hashlib
import io
import json
from datetime import UTC, datetime, timedelta
from uuid import uuid4

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from agent_workbench.config.settings import Settings
from agent_workbench.contracts.errors import ApiError
from agent_workbench.domain.attachment_parser import decode_text, parse_file
from agent_workbench.domain.attachments import (
    _parse_isolated,
    file_kind,
    parse_attachment,
    read_slice,
    validate_filename,
)
from agent_workbench.evidence.artifacts import ArtifactStore
from agent_workbench.persistence.tables import AttachmentRow


def pdf_fixture() -> bytes:
    writer = PdfWriter()
    for label in ("asset-one", "asset-two"):
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 10 100 Td ({label}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize(
    "filename",
    ["../secret.txt", "..\\secret.txt", "C:secret.txt", "bad\x00.txt", "a\n.txt", "a.", ""],
)
def test_filename_rejects_paths(filename):
    with pytest.raises(ApiError):
        validate_filename(filename)


@pytest.mark.parametrize(
    ("name", "raw"),
    [
        ("a.png", b"not png"),
        ("a.txt", b"MZbinary"),
        ("a.pdf", b"not pdf"),
        ("a.docx", b"not zip"),
        ("a.svg", b"<svg/>"),
    ],
)
def test_spoofed_and_forbidden_types(name, raw):
    with pytest.raises(ApiError):
        file_kind(name, raw)


def test_encodings_are_strict():
    assert decode_text("中文".encode())[0] == "中文"
    assert decode_text("中文".encode("utf-16"))[0] == "中文"
    with pytest.raises(UnicodeDecodeError):
        decode_text(b"\xffbad")
    with pytest.raises(ValueError):
        decode_text(b"abc\x00binary")


async def test_documents_and_image_in_isolated_process(tmp_path):
    cfg = Settings(data_root=tmp_path)
    pdf = tmp_path / "fixture.pdf"
    pdf.write_bytes(pdf_fixture())
    extracted = await parse_attachment(pdf, "pdf", cfg)
    assert extracted["status"] == "ready"
    assert extracted["metadata"]["page_count"] == 2
    assert "asset-two" in extracted["pages"][1]
    doc = Document()
    doc.add_heading("资产报告", 1)
    doc.add_paragraph("正文示例")
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "表格资产"
    docx = tmp_path / "fixture.docx"
    doc.save(str(docx))
    extracted = await parse_attachment(docx, "docx", cfg)
    assert extracted["status"] == "ready"
    assert all(
        word in extracted["pages"][0] for word in ("资产报告", "正文示例", "table", "表格资产")
    )
    image = tmp_path / "fixture.png"
    Image.new("RGB", (40, 20)).save(image)
    extracted = await parse_attachment(image, "png", cfg)
    assert extracted["status"] == "metadata_only"
    assert extracted["metadata"]["width"] == 40
    assert not extracted["pages"]
    malformed = tmp_path / "bad.pdf"
    malformed.write_bytes(b"%PDF-bad")
    assert (await parse_attachment(malformed, "pdf", cfg))["status"] == "failed"
    assert (await parse_attachment(pdf, "pdf", cfg.model_copy(update={"attachment_max_pages": 1})))[
        "status"
    ] == "failed"


def test_long_line_continuation_and_missing(tmp_path):
    cfg = Settings(data_root=tmp_path, attachment_read_chars=256)
    store = ArtifactStore(cfg)
    raw = ("中" * 600 + "\nlast\n").encode()
    original = store.write_bytes(raw, media_type="text/plain")
    path = store.resolve(original.relative_key)
    parsed = parse_file(path, "txt", 10000, 100)
    view = store.write_bytes(json.dumps(parsed).encode(), media_type="application/json")
    row = AttachmentRow(
        id=uuid4(),
        session_id=uuid4(),
        owner_id="local",
        upload_key="test",
        filename="fixture.txt",
        content_type="text/plain",
        byte_size=len(raw),
        sha256=hashlib.sha256(raw).hexdigest(),
        relative_key=original.relative_key,
        parsed_key=view.relative_key,
        parse_status="ready",
        metadata_json=parsed["metadata"],
        created_at=datetime.now(UTC),
        expires_at=datetime.now(UTC) + timedelta(hours=24),
        status="bound",
    )
    cursor = {}
    parts = []
    for _ in range(10):
        piece = read_slice(row, store, cfg, **cursor)
        assert len(piece.text) <= 256
        assert f"sha256:{row.sha256}" in piece.reference
        parts.append(piece.text)
        if not piece.next_cursor:
            break
        assert cursor != piece.next_cursor
        cursor = piece.next_cursor
    assert "last" in "".join(parts)
    assert sum(part.count("中") for part in parts) == 600
    assert read_slice(row, store, cfg, start_line=9999).status == "eof"
    path.unlink()
    assert read_slice(row, store, cfg).status == "missing"


def test_empty_opaque_and_parse_budget(tmp_path):
    empty = tmp_path / "empty"
    empty.write_bytes(b"")
    assert parse_file(empty, "txt", 1000, 20)["status"] == "empty"
    empty.write_bytes(b"MZopaque")
    result = parse_file(empty, "exe", 1000, 20)
    assert result["status"] == "unsupported" and not result["pages"]
    empty.write_bytes(b"a" * 1001)
    with pytest.raises(ValueError, match="limit"):
        parse_file(empty, "txt", 1000, 20)


@pytest.mark.parametrize("failure", ["timeout", "memory"])
def test_parser_resource_guard_kills_child_and_reports_failure(tmp_path, monkeypatch, failure):
    from types import SimpleNamespace

    class Child:
        pid = 123
        killed = False

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def poll(self):
            return 0 if self.killed else None

        def kill(self):
            self.killed = True

        def wait(self):
            return 0

    child = Child()
    monkeypatch.setattr(
        "agent_workbench.domain.attachments.subprocess.Popen", lambda *_a, **_k: child
    )
    monkeypatch.setattr(
        "agent_workbench.domain.attachments.psutil.Process",
        lambda _pid: SimpleNamespace(memory_info=lambda: SimpleNamespace(rss=1024 * 1024 * 1024)),
    )
    ticks = iter([0, 100 if failure == "timeout" else 0])
    monkeypatch.setattr("agent_workbench.domain.attachments.time.monotonic", lambda: next(ticks))
    result = _parse_isolated(tmp_path / "unused", "pdf", Settings(data_root=tmp_path))
    assert child.killed and result["status"] == "failed"
    assert result["metadata"]["error_type"] == (
        "TimeoutError" if failure == "timeout" else "MemoryError"
    )
    assert not result["pages"]
